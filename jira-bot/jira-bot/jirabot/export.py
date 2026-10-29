"""Export a reviewed grid to CSV or Excel (.xlsx, with the colour coding)."""
from __future__ import annotations

import csv
import io
from typing import Dict, List, Sequence

from .reviews import Mark


def _visible(columns: Sequence[str], hidden: Sequence[str]) -> List[str]:
    return [c for c in columns if c not in set(hidden)]


def _line(row: Dict[str, str], columns: Sequence[str]) -> List[str]:
    return [str(row.get(c, "")) for c in columns] + [row.get("_mark", ""), row.get("_note", "")]


def to_csv(columns: Sequence[str], hidden: Sequence[str], rows: Sequence[Dict[str, str]]) -> bytes:
    """CSV that opens cleanly in Excel (UTF-8 with a byte order mark)."""
    cols = _visible(columns, hidden)
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([c.title() if c != "key" else "Key" for c in cols] + ["Mark", "Note"])
    for row in rows:
        writer.writerow(_line(row, cols))
    return buffer.getvalue().encode("utf-8-sig")


def blend_with_white(hex_color: str, strength: float = 0.25) -> str:
    """A light version of a colour (Excel cannot do transparency)."""
    hex_color = hex_color.lstrip("#")
    r, g, b = (int(hex_color[i : i + 2], 16) for i in (0, 2, 4))
    mix = lambda c: int(round(255 - (255 - c) * strength))  # noqa: E731
    return f"{mix(r):02X}{mix(g):02X}{mix(b):02X}"


def xlsx_available() -> bool:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return False
    return True


def to_xlsx(
    columns: Sequence[str],
    hidden: Sequence[str],
    rows: Sequence[Dict[str, str]],
    marks: Sequence[Mark],
    sheet_name: str = "Review",
) -> bytes:
    """An .xlsx with a frozen header row, filters, sensible widths and rows tinted by mark."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cols = _visible(columns, hidden)
    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_name or "Review")[:31].replace("/", "-").replace("\\", "-")

    headers = [c.title() if c != "key" else "Key" for c in cols] + ["Mark", "Note"]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F2937")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    fills = {m.name: PatternFill("solid", fgColor=blend_with_white(m.hex)) for m in marks}
    for row in rows:
        ws.append(_line(row, cols))
        fill = fills.get(row.get("_mark", ""))
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = fill

    for index, header in enumerate(headers, start=1):
        longest = max([len(header)] + [len(str(ws.cell(r, index).value or "")) for r in range(2, min(ws.max_row, 300) + 1)])
        ws.column_dimensions[get_column_letter(index)].width = max(8, min(60, longest + 2))
    ws.freeze_panes = "A2"
    if ws.max_row >= 1:
        ws.auto_filter.ref = ws.dimensions

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
