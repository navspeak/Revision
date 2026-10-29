import io
import json
import os
import tempfile
import threading
import unittest
from pathlib import Path

import helpers  # noqa: F401
from helpers import make_settings

from jirabot import export
from jirabot.jira import AcliClient
from jirabot.reviews import DEFAULT_MARKS, Mark, ReviewStore, merge_rows

ROWS = [
    {"key": "ABC-1", "summary": "First", "status": "To Do", "priority": "High"},
    {"key": "ABC-2", "summary": "Second", "status": "In Progress", "priority": "Low"},
    {"key": "ABC-3", "summary": "Third", "status": "To Do", "priority": "Highest"},
]
COLS = ["key", "summary", "status", "priority"]


class ReviewStoreTests(unittest.TestCase):
    def setUp(self):
        self.store = ReviewStore(Path(tempfile.mkdtemp()) / "data" / "reviews.sqlite3")

    def test_snapshot_round_trip_and_defaults(self):
        n = self.store.save_snapshot("Sprint 42", "project = ABC", COLS, ROWS)
        self.assertEqual(n, 3)
        review = self.store.get("Sprint 42")
        self.assertEqual(review.jql, "project = ABC")
        self.assertEqual(review.columns, COLS)
        self.assertEqual(review.row_count, 3)
        self.assertEqual([m.name for m in review.marks], [m.name for m in DEFAULT_MARKS])
        self.assertEqual([r["key"] for r in self.store.load_rows("Sprint 42")], ["ABC-1", "ABC-2", "ABC-3"])

    def test_annotations_survive_a_rerun_of_the_query(self):
        self.store.save_snapshot("r", "q", COLS, ROWS)
        self.store.set_annotation("r", "ABC-2", mark="Move out", note="after the release")
        fresh = [dict(r, summary=r["summary"] + " (updated)") for r in ROWS[1:]] + [
            {"key": "ABC-9", "summary": "New", "status": "To Do", "priority": "Low"}
        ]
        self.store.save_snapshot("r", "q2", COLS, fresh)
        ann = self.store.annotations("r")["ABC-2"]
        self.assertEqual((ann.mark, ann.note), ("Move out", "after the release"))
        rows = self.store.load_rows("r")
        self.assertEqual([r["key"] for r in rows], ["ABC-2", "ABC-3", "ABC-9"])
        self.assertIn("(updated)", rows[0]["summary"])
        self.assertEqual(self.store.get("r").jql, "q2")

    def test_partial_updates_and_clearing(self):
        self.store.save_snapshot("r", "q", COLS, ROWS)
        self.store.set_annotation("r", "ABC-1", mark="Keep")
        self.store.set_annotation("r", "ABC-1", note="hello")  # note only: mark untouched
        ann = self.store.annotations("r")["ABC-1"]
        self.assertEqual((ann.mark, ann.note), ("Keep", "hello"))
        self.store.set_annotation("r", "ABC-1", mark="")
        self.assertEqual(self.store.annotations("r")["ABC-1"].mark, "")
        self.store.set_annotation("r", "ABC-1", note="")
        self.assertNotIn("ABC-1", self.store.annotations("r"))  # empty rows are removed

    def test_bulk_apply_and_counts(self):
        self.store.save_snapshot("r", "q", COLS, ROWS)
        self.assertEqual(self.store.apply_mark("r", ["ABC-1", "ABC-3"], "Not doing"), 2)
        self.store.set_annotation("r", "ABC-2", mark="Keep")
        self.assertEqual(self.store.mark_counts("r"), {"Not doing": 2, "Keep": 1})

    def test_reviews_are_independent(self):
        self.store.save_snapshot("a", "q", COLS, ROWS)
        self.store.save_snapshot("b", "q", COLS, ROWS)
        self.store.set_annotation("a", "ABC-1", mark="Keep")
        self.assertNotIn("ABC-1", self.store.annotations("b"))
        self.store.delete("a")
        self.assertIsNone(self.store.get("a"))
        self.assertEqual(len(self.store.load_rows("b")), 3)
        self.assertEqual(self.store.annotations("a"), {})

    def test_hidden_columns_and_custom_marks_persist(self):
        self.store.save_snapshot("r", "q", COLS, ROWS)
        self.store.set_hidden("r", ["priority"])
        self.store.set_marks("r", [Mark("Ship it", "green"), Mark("Park", "purple"), Mark("park", "red")])
        review = self.store.get("r")
        self.assertEqual(review.hidden, ["priority"])
        self.assertEqual([(m.name, m.color) for m in review.marks], [("Ship it", "green"), ("Park", "purple")])
        with self.assertRaises(ValueError):
            self.store.set_marks("r", [])

    def test_label_helpers(self):
        review = self.store.ensure("r")
        label = review.label_for("Move out")
        self.assertIn("Move out", label)
        self.assertNotEqual(label, "Move out")  # carries a colour square
        self.assertEqual(review.name_for_label(label), "Move out")
        self.assertEqual(review.name_for_label(""), "")

    def test_rows_without_keys_and_duplicates(self):
        n = self.store.save_snapshot("r", "q", ["summary"], [{"summary": "a"}, {"summary": "b"}, {"key": "X-1", "summary": "c"}, {"key": "X-1", "summary": "dup"}])
        self.assertEqual(n, 3)
        rows = self.store.load_rows("r")
        self.assertEqual(rows[0]["key"], "row-1")
        self.assertIn("key", self.store.get("r").columns)

    def test_names_are_cleaned(self):
        self.assertEqual(self.store.ensure("  My   review \n").name, "My review")
        with self.assertRaises(ValueError):
            self.store.ensure("   ")

    def test_large_review_is_fast_enough(self):
        import time
        rows = [{"key": f"ABC-{i}", "summary": f"Issue {i}", "status": "To Do", "priority": "Low"} for i in range(20000)]
        t0 = time.time()
        self.store.save_snapshot("big", "q", COLS, rows)
        self.store.apply_mark("big", [f"ABC-{i}" for i in range(0, 20000, 2)], "Keep")
        loaded = self.store.load_rows("big")
        merged = merge_rows(loaded, self.store.annotations("big"))
        self.assertEqual(len(merged), 20000)
        self.assertEqual(merged[0]["_mark"], "Keep")
        self.assertEqual(merged[1]["_mark"], "")
        self.assertLess(time.time() - t0, 15)

    def test_two_writers_do_not_corrupt(self):
        self.store.save_snapshot("r", "q", COLS, ROWS)
        other = ReviewStore(self.store.path)
        errors = []

        def work(store, mark):
            try:
                for _ in range(30):
                    store.set_annotation("r", "ABC-1", mark=mark, note=mark)
            except Exception as exc:  # noqa: BLE001
                errors.append(exc)

        threads = [threading.Thread(target=work, args=(s, m)) for s, m in ((self.store, "Keep"), (other, "Discuss"))]
        [t.start() for t in threads]
        [t.join() for t in threads]
        self.assertEqual(errors, [])
        self.assertIn(self.store.annotations("r")["ABC-1"].mark, ("Keep", "Discuss"))


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.rows = merge_rows(ROWS, {})
        self.rows[0]["_mark"], self.rows[0]["_note"] = "Move out", "after the release, then \"revisit\""
        self.rows[1]["_mark"] = "Not doing"
        self.marks = list(DEFAULT_MARKS)

    def test_csv_respects_hidden_columns_and_quotes(self):
        data = export.to_csv(COLS, ["priority"], self.rows)
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))  # BOM so Excel reads UTF-8
        import csv
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        self.assertEqual(rows[0], ["Key", "Summary", "Status", "Mark", "Note"])
        self.assertEqual(rows[1][3:], ["Move out", 'after the release, then "revisit"'])
        self.assertEqual(len(rows), 4)

    def test_xlsx_has_colours_filters_and_frozen_header(self):
        if not export.xlsx_available():
            self.skipTest("openpyxl not installed")
        import openpyxl
        data = export.to_xlsx(COLS, [], self.rows, self.marks, "Sprint 42 / triage")
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb.active
        self.assertEqual(ws.title, "Sprint 42 - triage")
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertIsNotNone(ws.auto_filter.ref)
        self.assertEqual([c.value for c in ws[1]], ["Key", "Summary", "Status", "Priority", "Mark", "Note"])
        orange = ws.cell(2, 1).fill.fgColor.rgb[-6:]
        red = ws.cell(3, 1).fill.fgColor.rgb[-6:]
        self.assertNotEqual(orange, red)
        self.assertEqual(ws.cell(4, 1).fill.fill_type, None)  # unmarked row has no colour
        self.assertEqual(ws.cell(2, 6).value, 'after the release, then "revisit"')

    def test_blend_with_white(self):
        self.assertEqual(export.blend_with_white("#000000", 0.5), "808080")
        self.assertEqual(export.blend_with_white("#ffffff", 0.2), "FFFFFF")


class PaginateTests(unittest.TestCase):
    def setUp(self):
        self.acli = AcliClient(make_settings())

    def test_paginate_fetches_everything(self):
        os.environ["FAKE_ACLI_TOTAL"] = "120"
        try:
            res = self.acli.search("project = ABC AND created >= -30d", paginate=True)
            self.assertEqual(len(res.rows), 120)
            capped = self.acli.search("project = ABC AND created >= -30d", limit=40)
            self.assertEqual(len(capped.rows), 40)
        finally:
            del os.environ["FAKE_ACLI_TOTAL"]

    def test_paginate_falls_back_to_limit_when_flag_unknown(self):
        os.environ["FAKE_ACLI_TOTAL"] = "120"
        os.environ["FAKE_ACLI_NOPAGINATE"] = "1"
        try:
            res = self.acli.search("project = ABC AND created >= -30d", limit=70, paginate=True)
            self.assertEqual(len(res.rows), 70)
        finally:
            del os.environ["FAKE_ACLI_TOTAL"]
            del os.environ["FAKE_ACLI_NOPAGINATE"]


if __name__ == "__main__":
    unittest.main()
